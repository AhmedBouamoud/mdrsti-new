# -*- coding: utf-8 -*-
"""
شرائح شرح بصرية لقالب الأستاذ أحمد بوعمود: خط زمني، بطاقات، سلسلة سببية، رباعية،
بطاقات مفاهيم، أسئلة مرقّمة، فكرة محورية، مهمة. عبارات قصيرة وخط كبير للعرض بالمسلاط.
"""
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

NAVY = RGBColor(0x0B, 0x22, 0x42)
NAVY2 = RGBColor(0x1C, 0x3A, 0x63)
GOLD = RGBColor(0xC9, 0xA2, 0x4A)
GOLD_LT = RGBColor(0xF3, 0xE6, 0xC4)
CREAM = RGBColor(0xF5, 0xF1, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x26, 0x30, 0x3D)
MUTED = RGBColor(0x5B, 0x66, 0x75)
ACC = [RGBColor(0x1F, 0x6F, 0x78), RGBColor(0x8C, 0x2F, 0x39), RGBColor(0x4F, 0x7A, 0x3A),
       RGBColor(0xB0, 0x6A, 0x1E), RGBColor(0x4A, 0x4F, 0x8C), RGBColor(0x0B, 0x22, 0x42)]
FONT = "Cairo"

# منطقة المحتوى في تخطيط «شرح» (10 × 5.625 بوصة)
X0, X1, Y0, Y1 = 0.45, 9.55, 1.2, 5.05


def _font(run, size, bold=False, color=INK):
    run.font.size = Pt(size); run.font.bold = bold; run.font.name = FONT
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr(); rPr.set("lang", "ar-MA")
    for tag in ("a:latin", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=20, bold=False, color=INK, align="r", anchor="t", spacing=1.0):
    """lines: نص أو قائمة (نص | (نص, حجم, عريض, لون))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        t, sz, b, c = (ln, size, bold, color) if isinstance(ln, str) else (list(ln) + [size, bold, color][len(ln) - 1:])[:4]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr(); pPr.set("rtl", "1"); pPr.set("algn", align)
        p.line_spacing = spacing
        r = p.add_run(); r.text = t; _font(r, sz, b, c)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if not shadow:
        sp = s._element.spPr
        ef = sp.makeelement(qn("a:effectLst"), {}); sp.append(ef)
    s.text_frame.text = ""
    return s


def circle(slide, cx, cy, d, fill, label=None, size=18, color=WHITE):
    s = box(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    if label is not None:
        tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(label); _font(r, size, True, color)
    return s


def arrow_left(slide, x, y, w, h, fill=GOLD):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.LEFT_ARROW)


def arrow_down(slide, x, y, w, h, fill=GOLD):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.DOWN_ARROW)


# ---------------- أنواع الشرائح ----------------
def timeline(slide, items):
    """خط زمني أفقي من اليمين إلى اليسار. item: {date, title, note}"""
    n = len(items)
    yl = 2.75
    seg = (X1 - X0) / n
    box(slide, X0 + 0.1, yl - 0.035, X1 - X0 - 0.2, 0.07, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    for i, it in enumerate(items):
        cx = X1 - seg * (i + 0.5)
        col = ACC[i % len(ACC)]
        # التاريخ فوق الخط
        pill = box(slide, cx - seg / 2 + 0.05, 1.35, seg - 0.1, 0.8, fill=col, radius=0.35)
        text(slide, cx - seg / 2 + 0.05, 1.35, seg - 0.1, 0.8, it["date"], 21, True, WHITE, "ctr", "m")
        box(slide, cx - 0.015, 2.15, 0.03, yl - 2.15, fill=col, shape=MSO_SHAPE.RECTANGLE)
        circle(slide, cx, yl, 0.3, WHITE)
        circle(slide, cx, yl, 0.2, col)
        # الحدث والدلالة تحت الخط
        text(slide, cx - seg / 2 + 0.02, yl + 0.25, seg - 0.04, 0.95, it["title"], 19, True, NAVY, "ctr")
        if it.get("note"):
            text(slide, cx - seg / 2 + 0.02, yl + 1.3, seg - 0.04, 1.0, it["note"], 17, False, MUTED, "ctr")


def cards(slide, items, cols=3):
    """بطاقات: item {title, text, tag?}"""
    n = len(items); rows = (n + cols - 1) // cols
    gap = 0.18
    cw = (X1 - X0 - gap * (cols - 1)) / cols
    ch = (Y1 - Y0 - gap * (rows - 1)) / rows
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = X1 - (c + 1) * cw - c * gap
        y = Y0 + r * (ch + gap)
        col = ACC[i % len(ACC)]
        box(slide, x, y, cw, ch, fill=WHITE, line=RGBColor(0xDD, 0xD5, 0xC4))
        box(slide, x + cw - 0.12, y, 0.12, ch, fill=col, shape=MSO_SHAPE.RECTANGLE)
        big1 = rows == 1
        tsz = 28 if big1 else 21
        th = (1.15 if max(len(t["title"]) for t in items) > 10 else 0.65) if big1 else 0.52
        ty0 = y + (0.45 if big1 else 0.08)
        text(slide, x + 0.15, ty0, cw - 0.4, th, it["title"], tsz, True, col, "r", "t", 0.95)
        ty = ty0 + th
        if it.get("tag"):
            text(slide, x + 0.15, ty, cw - 0.4, 0.34, it["tag"], 14, True, GOLD)
            ty += 0.34
        text(slide, x + 0.15, ty + 0.05, cw - 0.4, y + ch - ty - 0.1, it["text"], 24 if big1 else 18, False, INK, "r", "t", 1.0)


def quad(slide, items):
    """أربع لوحات ملوّنة 2×2: item {title, text}"""
    gap = 0.2
    w = (X1 - X0 - gap) / 2; h = (Y1 - Y0 - gap) / 2
    for i, it in enumerate(items[:4]):
        r, c = divmod(i, 2)
        x = X1 - (c + 1) * w - c * gap; y = Y0 + r * (h + gap)
        col = ACC[i]
        box(slide, x, y, w, h, fill=col)
        circle(slide, x + w - 0.45, y + 0.45, 0.55, WHITE, i + 1, 20, col)
        text(slide, x + 0.2, y + 0.15, w - 0.9, 0.65, it["title"], 30, True, WHITE)
        text(slide, x + 0.2, y + 0.9, w - 0.4, h - 1.0, it["text"], 22, False, WHITE)


def chain(slide, steps, caption=None):
    """سلسلة سببية متعرّجة من اليمين إلى اليسار (صفّان)."""
    n = len(steps); per = (n + 1) // 2
    bw, bh, aw = 2.35, 1.05, 0.55
    total = per * bw + (per - 1) * aw
    x_start = X1 - (X1 - X0 - total) / 2
    ys = [Y0 + 0.25, Y0 + 2.05]
    for i, st in enumerate(steps):
        row = i // per; k = i % per
        if row == 1:
            k = per - 1 - k  # الصف الثاني يعود من اليسار إلى اليمين
        x = x_start - (k + 1) * bw - k * aw
        y = ys[row]
        col = ACC[i % len(ACC)]
        box(slide, x, y, bw, bh, fill=col)
        text(slide, x + 0.06, y, bw - 0.12, bh, st, 20, True, WHITE, "ctr", "m", 0.95)
        if i < n - 1:
            nrow = (i + 1) // per
            if nrow == row:
                if row == 0:
                    arrow_left(slide, x - aw + 0.05, y + bh / 2 - 0.2, aw - 0.1, 0.4)
                else:
                    a = arrow_left(slide, x + bw + 0.05, y + bh / 2 - 0.2, aw - 0.1, 0.4); a.rotation = 180
            else:
                arrow_down(slide, x + bw / 2 - 0.2, y + bh + 0.1, 0.4, ys[1] - y - bh - 0.2)
    if caption:
        box(slide, X0, Y1 - 0.62, X1 - X0, 0.62, fill=GOLD_LT, radius=0.3)
        text(slide, X0 + 0.2, Y1 - 0.62, X1 - X0 - 0.4, 0.62, caption, 18, True, NAVY, "ctr", "m")


def questions(slide, items):
    n = len(items)
    h = (Y1 - Y0) / n
    for i, q in enumerate(items):
        y = Y0 + i * h
        circle(slide, X1 - 0.3, y + h / 2, 0.5, ACC[i % len(ACC)], i + 1, 20)
        box(slide, X0, y + 0.06, X1 - X0 - 0.7, h - 0.12, fill=WHITE, line=RGBColor(0xDD, 0xD5, 0xC4), radius=0.25)
        text(slide, X0 + 0.15, y + 0.06, X1 - X0 - 1.0, h - 0.12, q, 22, False, INK, "r", "m")


def concepts(slide, items):
    """بطاقات مفاهيم 3×2: item {term, def}"""
    cols, rows, gap = 3, (len(items) + 2) // 3, 0.16
    cw = (X1 - X0 - gap * (cols - 1)) / cols
    ch = (Y1 - Y0 - gap * (rows - 1)) / rows
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = X1 - (c + 1) * cw - c * gap; y = Y0 + r * (ch + gap)
        box(slide, x, y, cw, ch, fill=WHITE, line=RGBColor(0xDD, 0xD5, 0xC4), radius=0.08)
        hd = box(slide, x, y, cw, 0.58, fill=NAVY, radius=0.2)
        text(slide, x + 0.1, y, cw - 0.2, 0.58, it["term"], 22, True, WHITE, "ctr", "m")
        text(slide, x + 0.12, y + 0.62, cw - 0.24, ch - 0.68, it["def"], 20, False, INK, "ctr", "m", 1.0)


def big(slide, statement, sub=None):
    box(slide, X0 + 0.4, Y0 + 0.25, X1 - X0 - 0.8, Y1 - Y0 - 0.5, fill=WHITE, line=GOLD, radius=0.06)
    text(slide, X1 - 1.15, Y0 + 0.05, 0.8, 0.9, "”", 66, True, GOLD, "r")
    lines = [(statement, 38, True, NAVY)]
    if sub:
        lines.append((" ", 12, False, INK))
        lines.append((sub, 23, False, MUTED))
    text(slide, X0 + 0.9, Y0 + 0.5, X1 - X0 - 1.8, Y1 - Y0 - 1.0, lines, anchor="m", align="ctr", spacing=1.1)


def task(slide, steps, title="المهمة"):
    box(slide, X0, Y0, X1 - X0, Y1 - Y0, fill=WHITE, line=GOLD, radius=0.05)
    box(slide, X1 - 2.2, Y0 - 0.0, 2.2, 0.55, fill=GOLD, radius=0.3)
    text(slide, X1 - 2.2, Y0, 2.2, 0.55, title, 20, True, NAVY, "ctr", "m")
    h = (Y1 - Y0 - 0.8) / len(steps)
    for i, s in enumerate(steps):
        y = Y0 + 0.7 + i * h
        circle(slide, X1 - 0.5, y + h / 2, 0.45, NAVY, i + 1, 18)
        text(slide, X0 + 0.3, y, X1 - X0 - 1.2, h, s, 24, False, INK, "r", "m")


KINDS = {"timeline": timeline, "cards": cards, "quad": quad, "chain": chain,
         "questions": questions, "concepts": concepts, "big": big, "task": task}


def render(slide, spec):
    k = spec["type"]
    if k == "timeline": timeline(slide, spec["items"])
    elif k == "cards": cards(slide, spec["items"], spec.get("cols", 3))
    elif k == "quad": quad(slide, spec["items"])
    elif k == "chain": chain(slide, spec["steps"], spec.get("caption"))
    elif k == "questions": questions(slide, spec["items"])
    elif k == "concepts": concepts(slide, spec["items"])
    elif k == "big": big(slide, spec["text"], spec.get("sub"))
    elif k == "task": task(slide, spec["steps"], spec.get("label", "المهمة"))
