#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
مولّد عرض البوربوينت لدرس من «رحلة الأولى باك» انطلاقًا من قالب الأستاذ أحمد بوعمود.
- الشرح: من مقاطع الدرس في content/*.json
- «أكتب في دفتري»: حرفيًا من حقول notebook (نص ملف الدرس كما هو، دون تلخيص).
الاستعمال: python3 scripts/build_pptx.py <slug> <template.pptx> <out.pptx>
"""
import json, glob, os, re, sys, copy
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAVY = RGBColor(0x0B, 0x22, 0x42)
GOLD = RGBColor(0xC9, 0xA2, 0x4A)
FONT = "Cairo"
ORD = ["الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس", "السابع", "الثامن", "التاسع"]


def load(slug):
    for f in glob.glob(os.path.join(ROOT, "content", "**", "*.json"), recursive=True):
        d = json.load(open(f, encoding="utf-8"))
        if d["meta"].get("slug") == slug or os.path.basename(f).endswith(slug + ".json") or d["meta"]["id"] == slug:
            return d
    for f in glob.glob(os.path.join(ROOT, "content", "**", "*.json"), recursive=True):
        if slug in os.path.basename(f):
            return json.load(open(f, encoding="utf-8"))
    sys.exit("لم يُعثر على الدرس: " + slug)


# ---------- أدوات النص ----------
def rtl_para(p, size, bold=False, color=None, align=PP_ALIGN.RIGHT):
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    pPr.set("algn", "r" if align == PP_ALIGN.RIGHT else "ctr")
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:cs", "a:latin"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", FONT)
        rPr.set("lang", "ar-MA")


def set_text(ph, items, size, bullets=False, bold_first=False):
    """items: قائمة فقرات؛ عنصر يبدأ بـ '##' عنوان صغير."""
    tf = ph.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        is_h = it.startswith("##")
        txt = it[2:].strip() if is_h else it
        if bullets and not is_h:
            txt = "• " + txt
        # الكلمة الأولى قبل «:» بالخط العريض
        if bold_first and ":" in txt and not is_h:
            a, b = txt.split(":", 1)
            r1 = p.add_run(); r1.text = a + ":"
            r2 = p.add_run(); r2.text = b
            rtl_para(p, size)
            r1.font.bold = True
            r1.font.color.rgb = NAVY
        else:
            r = p.add_run(); r.text = txt
            rtl_para(p, size, bold=is_h, color=NAVY if is_h else None)
        p.space_after = Pt(6)
        # إلغاء ترقيم القالب حين نضيف «•» بأنفسنا
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("a:buNone")) is None:
            pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))


def fit(items, budget):
    """يقسّم الفقرات على شرائح حسب ميزانية الحروف، ويقسّم الفقرة الطويلة عند الجمل."""
    pieces = []
    for it in items:
        if len(it) <= budget:
            pieces.append(it)
            continue
        sents = re.split(r"(?<=[\.؛!؟])\s+", it)
        cur = ""
        for s in sents:
            if cur and len(cur) + len(s) + 1 > budget:
                pieces.append(cur); cur = s
            else:
                cur = (cur + " " + s).strip()
        if cur:
            pieces.append(cur)
    pages, cur, n = [], [], 0
    for it in pieces:
        w = len(it) + 25
        if cur and n + w > budget and not cur[-1].startswith("##"):
            pages.append(cur); cur, n = [], 0
        cur.append(it); n += w
    if cur:
        pages.append(cur)
    return pages


# ---------- الشرائح ----------
class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            self.prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
        self.L = {l.name: l for l in self.prs.slide_layouts}

    def phs(self, slide):
        return [p for p in slide.placeholders if p.placeholder_format.type not in (13,)]  # 13 = رقم الشريحة

    def cover(self, level, title, problem):
        s = self.prs.slides.add_slide(self.L["غلاف الدرس"])
        ph = self.phs(s)
        set_text(ph[0], [level], 20)
        set_text(ph[1], [title], 26 if len(title) > 45 else 36)
        set_text(ph[2], ["الإشكالية: " + problem], 14 if len(problem) > 160 else 18)

    def axis(self, n, title):
        s = self.prs.slides.add_slide(self.L["محور"])
        ph = self.phs(s)
        set_text(ph[0], [str(n)], 96)
        set_text(ph[1], [title], 36 if len(title) > 40 else 40)

    def explain(self, title, items, bullets=True, bold_first=True, budget=430):
        for i, page in enumerate(fit(items, budget)):
            s = self.prs.slides.add_slide(self.L["شرح"])
            ph = self.phs(s)
            set_text(ph[0], [title + (" (تابع)" if i else "")], 28 if len(title) > 38 else 32)
            size = 26 if sum(len(x) for x in page) < 380 else 24
            set_text(ph[1], page, size, bullets=bullets, bold_first=bold_first)

    def _nb_slide(self):
        s = self.prs.slides.add_slide(self.L["أكتب في دفتري"])
        return s, self.phs(s)[0]

    def notebook(self, blocks):
        """كل كتلة تُنسخ كما هي. الفقرات والقوائم بخط كبير، والجداول تُقسَّم على شرائح."""
        text_buf = []

        def flush():
            if not text_buf:
                return
            for page in fit(text_buf, 420):
                s, ph = self._nb_slide()
                size = 26 if sum(len(x) for x in page) < 300 else 24
                set_text(ph, page, size)
            text_buf.clear()

        heading = None
        for b in blocks:
            if "h" in b:
                heading = b["h"]
                text_buf.append("##" + b["h"])
            elif "p" in b:
                text_buf.append(b["p"])
            elif "ul" in b:
                text_buf.extend("• " + x for x in b["ul"])
            elif "table" in b:
                only_heading = len(text_buf) == 1 and text_buf[0].startswith("##")
                if not only_heading:
                    flush()
                else:
                    text_buf.clear()
                self.nb_table(b["table"], heading)
        flush()

    def nb_table(self, t, heading):
        ncol = len(t["head"])
        per = 3
        rows = t["rows"]
        for i in range(0, len(rows), per):
            chunk = rows[i:i + per]
            s, ph = self._nb_slide()
            L, T, W = Inches(0.6), Inches(1.3), Inches(8.5)   # منطقة الدفتر في القالب
            top = T
            if heading:
                set_text(ph, ["##" + heading + (" (تابع)" if i else "")], 22)
                ph.left, ph.top, ph.width, ph.height = L, T, W, Inches(0.55)
                top = T + Inches(0.6)
            else:
                ph._element.getparent().remove(ph._element)
            longest = max(len(c) for r in chunk for c in r)
            size = 20 if longest > 70 else 22
            gt = s.shapes.add_table(len(chunk) + 1, ncol, L, top,
                                    W, Inches(0.5 * (len(chunk) + 1))).table
            tblPr = gt._tbl.tblPr
            tblPr.set("rtl", "1")
            widths = [0.24, 0.26, 0.50] if ncol == 3 else ([0.28, 0.72] if ncol == 2 else [1 / ncol] * ncol)
            for j, w in enumerate(widths):
                gt.columns[j].width = int(W * w)
            for j, h in enumerate(t["head"]):
                c = gt.cell(0, j); c.text = h
                rtl_para(c.text_frame.paragraphs[0], size, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)
                c.fill.solid(); c.fill.fore_color.rgb = NAVY
            for r, row in enumerate(chunk, start=1):
                for j, val in enumerate(row):
                    c = gt.cell(r, j); c.text = val
                    c.vertical_anchor = MSO_ANCHOR.MIDDLE
                    rtl_para(c.text_frame.paragraphs[0], size, bold=(j == 0), color=NAVY if j == 0 else RGBColor(0x22, 0x22, 0x22))
                    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0xFD, 0xF6) if r % 2 else RGBColor(0xF6, 0xEE, 0xD8)

    def closing(self):
        s = self.prs.slides.add_slide(self.L["ختام"])
        set_text(self.phs(s)[0], ["شكراً لانتباهكم"], 40)


# ---------- استخراج الشرح من مقاطع الدرس ----------
def explain_items(sec):
    k = sec["kind"]
    if k == "table":
        return [" — ".join(r[:2]) + ": " + r[-1] if len(r) > 2 else r[0] + ": " + r[1] for r in sec["rows"]] + \
               ([sec["note"]["label"] + ": " + sec["note"]["text"]] if sec.get("note") else [])
    if k == "cards":
        out = []
        for c in sec["cards"]:
            body = "؛ ".join(c["bullets"]) if c.get("bullets") else c.get("body", "")
            out.append(c["cap"] + ": " + body)
        return out
    if k == "questions":
        return list(sec["items"])
    if k == "text":
        return list(sec.get("paragraphs", [])) + ([sec["note"]["label"] + ": " + sec["note"]["text"]] if sec.get("note") else [])
    if k == "application":
        return [sec["task"]]
    return []


def short_title(t):
    return re.sub(r"^[^:]{1,12}:\s*", "", t).strip()


def build(slug, template, out):
    d = load(slug)
    m = d["meta"]
    deck = Deck(template)
    deck.cover(f'الأولى باكالوريا | {m["subject"]}', m["title"], d["problematic"])
    deck.explain("مقدمة", [d["intro"]], bullets=False, bold_first=False)
    if d.get("notebook_intro"):
        deck.notebook(d["notebook_intro"]["blocks"])
    pending, n = [], 0
    for sec in d["sections"]:
        if sec["kind"] in ("resources", "infographic"):
            continue
        pending.append(sec)
        if sec.get("notebook"):
            n += 1
            deck.axis(n, f"المحور {ORD[n-1]}: " + short_title(sec["title"]))
            for p in pending:
                deck.explain(short_title(p["title"]), explain_items(p),
                             bullets=p["kind"] != "text", bold_first=p["kind"] in ("cards", "table"))
            deck.notebook(sec["notebook"]["blocks"])
            pending = []
    for p in pending:  # ما تبقّى (مثل التطبيق المنهجي)
        deck.explain(short_title(p["title"]), explain_items(p), bullets=False, bold_first=False)
    deck.explain("خاتمة", [d["conclusion"]], bullets=False, bold_first=False)
    if d.get("notebook_conclusion"):
        deck.notebook(d["notebook_conclusion"]["blocks"])
    deck.closing()
    deck.prs.save(out)
    print("✅", out, "—", len(deck.prs.slides), "شريحة")


if __name__ == "__main__":
    build(*sys.argv[1:4])
